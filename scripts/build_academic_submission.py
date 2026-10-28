"""Build the academic DOCX from the reviewed Markdown and original evidence images."""

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/INFORME_ACADEMICO.md"
OUTPUT = ROOT / "output/documents/INFORME_ACADEMICO_FINAL.docx"
WIDTH = 6.2677


def inline(paragraph, text):
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*|\[[^\]]+\]\([^)]+\))", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        italic = part.startswith("*") and part.endswith("*") and not bold
        code = part.startswith("`") and part.endswith("`")
        if bold:
            value = part[2:-2]
        elif italic or code:
            value = part[1:-1]
        elif part.startswith("["):
            value = re.match(r"\[([^\]]+)\]", part).group(1)
        else:
            value = part
        value = value.replace("\u2011", "-").replace("\u2013", "-").replace("\u2014", "-")
        run = paragraph.add_run(value)
        run.bold = bold
        run.italic = italic
        if code:
            run.font.name = "Consolas"
            run.font.size = Pt(9)


def widths_for(index, n):
    preferred = {
        1: [0.62, 2.55, WIDTH - 3.17],
        2: [1.48, 2.42, WIDTH - 3.9],
        3: [0.95, WIDTH - 0.95],
        4: [0.9, 2.45, WIDTH - 3.35],
        5: [1.05, 1.93, 2.25, WIDTH - 5.23],
        6: [4.28, 1.04, WIDTH - 5.32],
        7: [2.05, 2.1, WIDTH - 4.15],
        8: [0.90, 0.85, 0.82, 0.72, 1.40, WIDTH - 4.69],
        9: [1.32, 2.45, WIDTH - 3.77],
        10: [0.6, 2.82, WIDTH - 3.42],
    }
    return preferred.get(index, [WIDTH / n] * n)


def set_table(doc, data, index):
    table = doc.add_table(rows=1, cols=len(data[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = widths_for(index, len(data[0]))
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + edge)
        for key, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
            border.set(qn("w:" + key), value)
        borders.append(border)
    pr.append(borders)
    for row_number, values in enumerate(data):
        row = table.rows[0] if row_number == 0 else table.add_row()
        tr_pr = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        tr_pr.append(no_split)
        if row_number == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))
        for cell, value, width in zip(row.cells, values, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E7E6E6" if row_number == 0 else "FFFFFF")
            cell._tc.get_or_add_tcPr().append(shade)
            margins = OxmlElement("w:tcMar")
            for side in ("top", "bottom", "left", "right"):
                margin = OxmlElement("w:" + side)
                margin.set(qn("w:w"), "85")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            cell._tc.get_or_add_tcPr().append(margins)
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.space_before = Pt(3)
            inline(p, value)
            for run in p.runs:
                run.font.size = Pt(10)
                if row_number == 0:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.2677), Inches(11.6929)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
        st = doc.styles[name]
        st.font.name = "Times New Roman"
        st.font.color.rgb = RGBColor(0, 0, 0)
        fonts = st.element.get_or_add_rPr().get_or_add_rFonts()
        for key in list(fonts.attrib):
            if key.endswith("Theme"):
                del fonts.attrib[key]
        fonts.set(qn("w:ascii"), "Times New Roman")
        fonts.set(qn("w:hAnsi"), "Times New Roman")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Inches(0.5)
    normal.paragraph_format.widow_control = True
    doc.styles["Title"].font.size = Pt(22)
    doc.styles["Subtitle"].font.size = Pt(13)
    for name, size in [("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 12)]:
        st = doc.styles[name]
        st.font.size = Pt(size)
        st.font.bold = True
        st.paragraph_format.first_line_indent = Inches(0)
        st.paragraph_format.line_spacing = 1.2
        st.paragraph_format.space_before = Pt(16)
        st.paragraph_format.space_after = Pt(8)
    for name in ["Code", "Evidence Note"]:
        if name not in doc.styles:
            doc.styles.add_style(name, 1)
        st = doc.styles[name]
        st.font.name = "Consolas" if name == "Code" else "Times New Roman"
        st.font.size = Pt(8.5 if name == "Code" else 10.5)
        st.paragraph_format.line_spacing = 1.05 if name == "Code" else 1.15
        st.paragraph_format.first_line_indent = Inches(0)
        st.paragraph_format.space_after = Pt(8)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    header._p.append(field)
    doc.add_paragraph("Informe académico de LogisTech", "Title")
    doc.add_paragraph("Versión 2.0 y evidencias de pruebas del núcleo inicial", "Subtitle")
    p = doc.add_paragraph("Desarrollo de Soluciones en la Nube   TECSUP")
    p.paragraph_format.first_line_indent = Inches(0)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 1
    tables = 0
    figures = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            p = doc.add_paragraph("\n".join(block), "Code")
            p.paragraph_format.keep_together = True
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in row):
                    rows.append(row)
                i += 1
            tables += 1
            set_table(doc, rows, tables)
            continue
        image = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        if image:
            path = SOURCE.parent / image[2]
            with Image.open(path) as im:
                w, h = im.size
            width = min(WIDTH, 5.7 * w / h)
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.line_spacing = 1
            p.paragraph_format.keep_with_next = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic = p.add_run().add_picture(str(path), width=Inches(width))
            pic._inline.docPr.set("descr", image[1])
            figures += 1
            i += 1
            continue
        heading = re.match(r"(#{2,4}) (.+)", line)
        if heading:
            if heading[2].startswith(("Anexo A", "10. Conclusiones")):
                doc.add_page_break()
            doc.add_heading(heading[2], level=len(heading[1]) - 1)
            i += 1
            continue
        style = "Evidence Note" if line.startswith("*Nota.*") else None
        if re.match(r"^\d+\. ", line):
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_together = True
            inline(p, line)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.line_spacing = 1.3
            inline(p, line[2:])
        else:
            text = line
            i_next = i + 1
            while (
                i_next < len(lines)
                and lines[i_next].strip()
                and not re.match(r"^(#|\||```|!\[|- |\d+\. )", lines[i_next])
            ):
                text += " " + lines[i_next].strip()
                i_next += 1
            i = i_next - 1
            p = doc.add_paragraph(style=style)
            inline(p, text)
            if re.fullmatch(r"\*\*(Tabla|Figura) \d+\*\*", text):
                p.paragraph_format.first_line_indent = Inches(0)
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.line_spacing = 1.2
            elif re.fullmatch(r"\*[^*]+\*", text):
                p.paragraph_format.first_line_indent = Inches(0)
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.line_spacing = 1.2
                p.paragraph_format.space_after = Pt(6)
        i += 1
    # Remove template title rules and theme residue; preserve intentional table borders.
    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    for paragraph in doc.paragraphs:
        for border in list(paragraph._p.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    props = doc.core_properties
    props.title = "Informe académico de LogisTech 2.0 y evidencias de pruebas"
    props.subject = "Núcleo inicial F0–F5 y pruebas locales documentadas"
    props.author = ""
    props.last_modified_by = ""
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"DOCX created: {OUTPUT}; tables={tables}; embedded figures={figures}")


if __name__ == "__main__":
    main()
